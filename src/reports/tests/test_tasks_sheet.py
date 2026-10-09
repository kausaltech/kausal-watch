"""The tasks sheet of the report export."""

from __future__ import annotations

import datetime
import io
from typing import Any

from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.utils.translation import gettext as _
from reversion.models import Revision, Version
from reversion.revisions import add_to_revision, create_revision

import pytest
from openpyxl import load_workbook

from actions.models.action import ActionTask, ActionTaskContactPerson, ActionTaskResponsibleParty
from actions.models.features import PlanFeatures
from actions.tests.factories import ActionFactory, ActionTaskFactory, PlanFactory
from orgs.tests.factories import OrganizationFactory
from people.tests.factories import PersonFactory
from reports.export import export_dashboard_report_for_plan
from reports.tests.factories import ReportFactory, ReportTypeFactory

from .fixtures import *

pytestmark = pytest.mark.django_db


def _rows(report, user) -> dict[str, dict[str, Any]]:
    df = report.get_xlsx_exporter(user=user).generate_tasks_dataframe()
    return {row[_('Name')]: row for row in df.iter_rows(named=True)}


@pytest.fixture
def report(plan):
    # The report period starts 2023-12-15
    return ReportFactory.create(type=ReportTypeFactory.create(plan=plan))


def test_live_tasks_of_visible_actions(plan, action, report):
    ActionTaskFactory.create(
        action=action,
        name='Insulate the depot',
        due_at=datetime.date(2027, 1, 1),
        state=ActionTask.IN_PROGRESS,
        details='<p>Start with the <b>roof</b></p>',
    )
    internal_action = ActionFactory.create(plan=plan, visibility='internal')
    ActionTaskFactory.create(action=internal_action, name='Hidden task')

    rows = _rows(report, None)

    assert set(rows) == {'Insulate the depot'}
    row = rows['Insulate the depot']
    assert row[_('Action identifier')] == action.identifier
    assert row[_('Due date')] == datetime.date(2027, 1, 1)
    assert row[_('State')] == str(dict(ActionTask.STATES)[ActionTask.IN_PROGRESS])
    assert row[_('Details')] == 'Start with the **roof**'


def test_completed_action_keeps_task_state_at_completion(action, superuser, report):
    task = ActionTaskFactory.create(action=action, name='Insulate the depot', state=ActionTask.NOT_STARTED)
    action.mark_as_complete_for_report(report, superuser)

    task.state = ActionTask.COMPLETED
    task.completed_at = datetime.date(2024, 3, 1)
    task.save()

    row = _rows(report, superuser)['Insulate the depot']
    assert row[_('State')] == str(dict(ActionTask.STATES)[ActionTask.NOT_STARTED])
    assert row[_('Completion date')] is None


def test_state_at_start_of_reporting_period_comes_from_history(action, superuser, report):
    task = ActionTaskFactory.create(action=action, name='Insulate the depot', state=ActionTask.IN_PROGRESS)
    with create_revision():
        add_to_revision(action)
    Revision.objects.update(date_created=datetime.datetime(2023, 12, 1, tzinfo=datetime.UTC))
    task.state = ActionTask.COMPLETED
    task.completed_at = datetime.date(2024, 3, 1)
    task.save()
    ActionTaskFactory.create(action=action, name='Added later')

    rows = _rows(report, superuser)

    assert rows['Insulate the depot'][_('State')] == str(dict(ActionTask.STATES)[ActionTask.COMPLETED])
    assert rows['Insulate the depot'][_('State at start of reporting period')] == str(
        dict(ActionTask.STATES)[ActionTask.IN_PROGRESS]
    )
    assert rows['Added later'][_('State at start of reporting period')] == ''


def test_assignee_columns_follow_the_plan_feature(plan, action, superuser, report, monkeypatch):
    task = ActionTaskFactory.create(action=action, name='Insulate the depot')
    organization = OrganizationFactory.create(name='Department of Heat')
    plan.related_organizations.add(organization)
    person = PersonFactory.create(organization=plan.organization, first_name='Ada', last_name='Byron')
    ActionTaskResponsibleParty.objects.create(task=task, organization=organization)
    ActionTaskContactPerson.objects.create(task=task, person=person)

    assert _('Responsible organizations') not in _rows(report, superuser)['Insulate the depot']

    monkeypatch.setattr(PlanFeatures._meta.get_field('has_action_task_assignees'), 'default', True)
    plan.features.has_action_task_assignees = True
    plan.features.save()
    row = _rows(report, superuser)['Insulate the depot']
    assert row[_('Responsible organizations')] == 'Department of Heat'
    assert row[_('Contact persons')] == 'Ada Byron'


def test_sheet_uses_the_plan_term_and_is_only_in_report_export(plan_with_pages, superuser):
    report = ReportFactory.create(type=ReportTypeFactory.create(plan=plan_with_pages))
    ActionTaskFactory.create(action=ActionFactory.create(plan=plan_with_pages))
    sheet_name = str(plan_with_pages.general_content.get_action_task_term_display_plural())

    report.mark_as_complete(superuser)
    report_output = report.get_xlsx_exporter(user=superuser).generate_xlsx()
    dashboard_output, _filename = export_dashboard_report_for_plan(plan_with_pages, 'xlsx', superuser)
    assert isinstance(dashboard_output, bytes)

    assert sheet_name in load_workbook(io.BytesIO(report_output), read_only=True).sheetnames
    assert sheet_name not in load_workbook(io.BytesIO(dashboard_output), read_only=True).sheetnames


def test_queries_do_not_grow_with_tasks(plan, superuser, report):
    def count_queries(task_count: int) -> int:
        action = ActionFactory.create(plan=plan)
        for _i in range(task_count):
            ActionTaskFactory.create(action=action)
        exporter = report.get_xlsx_exporter(user=superuser)
        exporter._get_serialized_report_data()
        with CaptureQueriesContext(connection) as queries:
            exporter.generate_tasks_dataframe()
        return len(queries)

    assert count_queries(2) == count_queries(6)


def test_umbrella_report_names_the_plan_of_each_task(plan_with_pages, superuser):
    action_list_page = plan_with_pages.get_action_list_page()
    action_list_page.include_related_plans = True
    action_list_page.save()
    child = PlanFactory.create(parent=plan_with_pages, name='Child plan')
    ActionTaskFactory.create(action=ActionFactory.create(plan=plan_with_pages, identifier='1'), name='Parent task')
    ActionTaskFactory.create(action=ActionFactory.create(plan=child, identifier='1'), name='Child task')
    report = ReportFactory.create(type=ReportTypeFactory.create(plan=plan_with_pages))

    rows = _rows(report, superuser)

    assert rows['Parent task'][_('Plan')] == plan_with_pages.name
    assert rows['Child task'][_('Plan')] == 'Child plan'


def test_single_plan_report_has_no_plan_column(action, superuser, report):
    ActionTaskFactory.create(action=action, name='Insulate the depot')

    assert _('Plan') not in _rows(report, superuser)['Insulate the depot']


def test_details_saved_before_the_field_rename_are_shown(action, superuser, report):
    task = ActionTaskFactory.create(action=action, name='Insulate the depot', details='<p>Old text</p>')
    action.mark_as_complete_for_report(report, superuser)
    version = Version.objects.get_for_object(task).get()
    version.serialized_data = version.serialized_data.replace('"details":', '"comment":')
    version.save()

    assert _rows(report, superuser)['Insulate the depot'][_('Details')] == 'Old text'
