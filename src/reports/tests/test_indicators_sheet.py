"""The indicators sheet of the report export, and the freezing of indicator data with action and report completion."""

from __future__ import annotations

import datetime
import io
from typing import Any

from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.utils.translation import gettext as _, pgettext

import pytest
from openpyxl import load_workbook

from actions.tests.factories import ActionFactory, PlanFactory
from indicators.tests.factories import (
    ActionIndicatorFactory,
    DimensionCategoryFactory,
    IndicatorFactory,
    IndicatorGoalFactory,
    IndicatorValueFactory,
)
from orgs.tests.factories import OrganizationFactory
from reports.export import export_dashboard_report_for_plan
from reports.models import IndicatorSnapshot

from .fixtures import *

pytestmark = pytest.mark.django_db


def _rows(report, user) -> dict[str, dict[str, Any]]:
    df = report.get_xlsx_exporter(user=user).generate_indicators_dataframe()
    return {row[_('Indicator')]: row for row in df.iter_rows(named=True)}


@pytest.fixture
def report(plan, report_factory, report_type_factory):
    # The report period starts 2023-12-15
    return report_factory(type=report_type_factory(plan=plan))


@pytest.fixture
def indicator(plan, action):
    indicator = IndicatorFactory.create(plans=[plan], name='Emissions')
    ActionIndicatorFactory.create(action=action, indicator=indicator)
    return indicator


def test_live_data(plan, action, superuser, report, indicator):
    IndicatorValueFactory.create(indicator=indicator, date=datetime.date(2022, 12, 31), value=10.0)
    IndicatorValueFactory.create(indicator=indicator, date=datetime.date(2023, 12, 31), value=8.0)
    IndicatorValueFactory.create(indicator=indicator, date=datetime.date(2024, 12, 31), value=6.0)
    IndicatorValueFactory.create(
        indicator=indicator, date=datetime.date(2025, 12, 31), value=99.0, categories=[DimensionCategoryFactory.create()]
    )
    IndicatorGoalFactory.create(indicator=indicator, date=datetime.date(2030, 12, 31), value=2.0)
    IndicatorGoalFactory.create(indicator=indicator, date=datetime.date(2035, 12, 31), value=0.0)
    IndicatorFactory.create(plans=[plan], name='Unlinked')

    rows = _rows(report, superuser)

    row = rows['Emissions']
    assert row[_('Unit')] == str(indicator.unit)
    assert row[_('Organization')] == str(indicator.organization)
    assert row[_('Value at start of reporting period')] == 10.0
    assert row[_('Date of value at start of reporting period')] == datetime.date(2022, 12, 31)
    assert row[_('Latest value')] == 6.0
    assert row[_('Date of latest value')] == datetime.date(2024, 12, 31)
    assert row[_('Target value')] == 0.0
    assert row[_('Target date')] == datetime.date(2035, 12, 31)
    assert action.identifier in row[pgettext('Action model', 'Actions')]
    assert rows['Unlinked'][pgettext('Action model', 'Actions')] == ''
    assert _('Data frozen at') not in next(iter(rows.values()))


def test_action_completion_freezes_its_indicators(plan, action, superuser, report, indicator):
    value = IndicatorValueFactory.create(indicator=indicator, date=datetime.date(2024, 12, 31), value=6.0)
    unlinked = IndicatorFactory.create(plans=[plan], name='Unlinked')
    unlinked_value = IndicatorValueFactory.create(indicator=unlinked, date=datetime.date(2024, 12, 31), value=1.0)

    action.mark_as_complete_for_report(report, superuser)
    value.value = 7.0
    value.save()
    unlinked_value.value = 2.0
    unlinked_value.save()

    rows = _rows(report, superuser)
    assert rows['Emissions'][_('Latest value')] == 6.0
    assert rows['Emissions'][_('Data frozen at')] is not None
    assert rows['Unlinked'][_('Latest value')] == 2.0
    assert rows['Unlinked'][_('Data frozen at')] is None


def test_latest_action_completion_wins(plan, action, superuser, report, indicator):
    other_action = ActionFactory.create(plan=plan)
    ActionIndicatorFactory.create(action=other_action, indicator=indicator)
    value = IndicatorValueFactory.create(indicator=indicator, date=datetime.date(2024, 12, 31), value=6.0)

    action.mark_as_complete_for_report(report, superuser)
    value.value = 7.0
    value.save()
    other_action.mark_as_complete_for_report(report, superuser)
    value.value = 8.0
    value.save()

    assert _rows(report, superuser)['Emissions'][_('Latest value')] == 7.0


def test_undoing_latest_completion_restores_previous_snapshot(plan, action, superuser, report, indicator):
    other_action = ActionFactory.create(plan=plan)
    ActionIndicatorFactory.create(action=other_action, indicator=indicator)
    value = IndicatorValueFactory.create(indicator=indicator, date=datetime.date(2024, 12, 31), value=6.0)

    action.mark_as_complete_for_report(report, superuser)
    value.value = 7.0
    value.save()
    other_action.mark_as_complete_for_report(report, superuser)
    other_action.undo_marking_as_complete_for_report(report, superuser)

    assert _rows(report, superuser)['Emissions'][_('Latest value')] == 6.0


def test_deleting_completed_action_releases_its_snapshots(action, superuser, report, indicator):
    action.mark_as_complete_for_report(report, superuser)

    action.delete()

    assert not report.indicator_snapshots.exists()


def test_completed_report_keeps_indicators_frozen_by_a_deleted_action(action, superuser, report, indicator):
    value = IndicatorValueFactory.create(indicator=indicator, date=datetime.date(2024, 12, 31), value=6.0)
    action.mark_as_complete_for_report(report, superuser)
    value.value = 7.0
    value.save()
    report.mark_as_complete(superuser)

    action.delete()

    assert _rows(report, superuser)['Emissions'][_('Latest value')] == 6.0


def test_undoing_action_completion_keeps_snapshots_of_other_complete_actions(plan, action, superuser, report, indicator):
    other_action = ActionFactory.create(plan=plan)
    ActionIndicatorFactory.create(action=other_action, indicator=indicator)
    only_linked_to_action = IndicatorFactory.create(plans=[plan])
    ActionIndicatorFactory.create(action=action, indicator=only_linked_to_action)

    action.mark_as_complete_for_report(report, superuser)
    other_action.mark_as_complete_for_report(report, superuser)
    action.undo_marking_as_complete_for_report(report, superuser)

    assert set(report.indicator_snapshots.values_list('indicator', flat=True)) == {indicator.pk}


def test_undoing_action_completion_releases_indicators_unlinked_since(action, superuser, report, indicator):
    action.mark_as_complete_for_report(report, superuser)
    action.related_indicators.all().delete()

    action.undo_marking_as_complete_for_report(report, superuser)

    assert not report.indicator_snapshots.exists()


def test_report_completion_freezes_remaining_indicators(plan, action, superuser, report, indicator):
    unlinked = IndicatorFactory.create(plans=[plan], name='Unlinked')
    value = IndicatorValueFactory.create(indicator=unlinked, date=datetime.date(2024, 12, 31), value=1.0)
    IndicatorFactory.create(name='Other plan', plans=[PlanFactory.create()])
    action.mark_as_complete_for_report(report, superuser)

    report.mark_as_complete(superuser)
    value.value = 2.0
    value.save()
    IndicatorFactory.create(plans=[plan], name='Added after completion')

    rows = _rows(report, superuser)
    assert set(rows) == {'Emissions', 'Unlinked'}
    assert rows['Unlinked'][_('Latest value')] == 1.0

    report.undo_marking_as_complete(superuser)
    assert list(IndicatorSnapshot.objects.filter(report=report).values_list('indicator', flat=True)) == [indicator.pk]


@pytest.mark.usefixtures('indicator')
def test_internal_indicators_are_hidden_from_anonymous_users(plan, superuser, report):
    IndicatorFactory.create(plans=[plan], name='Internal', visibility='internal')

    assert set(_rows(report, None)) == {'Emissions'}
    assert set(_rows(report, superuser)) == {'Emissions', 'Internal'}


@pytest.mark.usefixtures('indicator')
def test_sheet_is_only_in_report_export(plan_with_pages, superuser, report):
    report_output = report.get_xlsx_exporter(user=superuser).generate_xlsx()
    dashboard_output, _filename = export_dashboard_report_for_plan(plan_with_pages, 'xlsx', superuser)
    assert isinstance(dashboard_output, bytes)

    assert _('Indicators') in load_workbook(io.BytesIO(report_output), read_only=True).sheetnames
    assert _('Indicators') not in load_workbook(io.BytesIO(dashboard_output), read_only=True).sheetnames


@pytest.mark.usefixtures('indicator')
def test_sheet_is_omitted_from_reports_completed_without_indicator_snapshots(superuser, report):
    report.mark_as_complete(superuser)
    report.indicator_snapshots.all().delete()

    output = report.get_xlsx_exporter(user=superuser).generate_xlsx()

    assert _('Indicators') not in load_workbook(io.BytesIO(output), read_only=True).sheetnames


@pytest.fixture
def umbrella_plan(plan_with_pages):
    action_list_page = plan_with_pages.get_action_list_page()
    action_list_page.include_related_plans = True
    action_list_page.save()
    return plan_with_pages


def test_indicators_of_hidden_child_plans_are_not_listed(umbrella_plan, superuser, report):
    hidden_child = PlanFactory.create(parent=umbrella_plan, published_at=None, visibility='internal')
    IndicatorFactory.create(plans=[umbrella_plan], name='Parent')
    IndicatorFactory.create(plans=[hidden_child], name='Hidden child')

    assert set(_rows(report, None)) == {'Parent'}
    assert set(_rows(report, superuser)) == {'Parent', 'Hidden child'}


def test_report_completion_freezes_only_the_report_plan_indicators(umbrella_plan, superuser, report):
    """Child plan indicators are left out of a completed report, like child plan actions."""
    child = PlanFactory.create(parent=umbrella_plan)
    IndicatorFactory.create(plans=[umbrella_plan], name='Parent')
    IndicatorFactory.create(plans=[child], name='Child')

    report.mark_as_complete(superuser)

    assert set(_rows(report, superuser)) == {'Parent'}


def test_snapshots_store_only_actions_of_the_report_plan(umbrella_plan, superuser, report):
    child = PlanFactory.create(parent=umbrella_plan)
    parent_action = ActionFactory.create(plan=umbrella_plan, identifier='P1')
    child_action = ActionFactory.create(plan=child, identifier='C1')
    indicator = IndicatorFactory.create(plans=[umbrella_plan], name='Shared')
    ActionIndicatorFactory.create(action=parent_action, indicator=indicator)
    ActionIndicatorFactory.create(action=child_action, indicator=indicator)
    actions_label = pgettext('Action model', 'Actions')

    assert _rows(report, superuser)['Shared'][actions_label] == 'P1, C1'

    parent_action.mark_as_complete_for_report(report, superuser)

    snapshot = report.indicator_snapshots.get(indicator=indicator)
    assert [a['identifier'] for a in snapshot.data['linked_actions']] == ['P1']
    assert _rows(report, superuser)['Shared'][actions_label] == 'P1'


def test_indicators_of_a_plan_hidden_from_the_user_are_not_listed(plan, report, indicator):
    plan.published_at = None
    plan.visibility = 'internal'
    plan.save()

    assert _rows(report, None) == {}


def test_report_completion_queries_do_not_grow_with_indicators(plan, superuser, report_factory, report_type_factory):
    parent_organization = OrganizationFactory.create()

    def count_queries(indicator_count: int) -> int:
        for _i in range(indicator_count):
            organization = OrganizationFactory.create(parent=parent_organization)
            IndicatorFactory.create(plans=[plan], organization=organization)
        report = report_factory(type=report_type_factory(plan=plan))
        with CaptureQueriesContext(connection) as queries:
            report.mark_as_complete(superuser)
        return len(queries)

    assert count_queries(2) == count_queries(6)
    snapshot = IndicatorSnapshot.objects.filter(indicator__plans=plan).select_related('indicator__organization').first()
    assert snapshot is not None
    assert snapshot.data['organization'] == str(snapshot.indicator.organization)
